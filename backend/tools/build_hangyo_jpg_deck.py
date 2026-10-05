"""Designer-vs-automated deck from two JPG folders (pairs.json made by matching file names).
    python tools/build_hangyo_jpg_deck.py <pairs.json> <out.pptx> [TITLE] [source note] [changes.json]
Any brand: pairs.json = [{d: designer image, a: automated image, name, sim, dsz, asz}, ...]
Several brands in ONE deck: pairs.json = {"Hangyo": [...], "Agarpathi": [...]} - a title slide, then for each brand an
at-a-glance slide and one slide per board.
changes.json (optional): {"slides": [{"title": "...", "subtitle": "...", "sections": [{"heading": "...", "items": ["...", ...]}]}]} - each entry
becomes a text slide right after the title slide (what was fixed in this round, what is still open)."""
import json, re, sys
from pptx import Presentation
from pptx.util import Inches, Pt
sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parents[1]))
from tools.build_comparison_deck import (tb, pill, card, place, bg, title, INK, RED, GREEN, AMBER, GREY, WHITE, SW, SH)
from pptx.dml.color import RGBColor

raw = json.load(open(sys.argv[1], encoding="utf-8"))
groups = raw if isinstance(raw, dict) else {None: raw}
TITLE = sys.argv[3] if len(sys.argv) > 3 else ("AGARPATHI + HANGYO" if isinstance(raw, dict) else "HANGYO")
NOTE = sys.argv[4] if len(sys.argv) > 4 else "designer JPEG vs automated picture from the app"
def ratio(sz): return sz[0] / sz[1]
def verdict(r):
    if not r["a"]: return "MISSING", GREY
    if not r.get("d"): return "NO DESIGNER JPEG", GREY
    asp = abs(ratio(r["dsz"]) / ratio(r["asz"]) - 1)
    if asp > 0.05: return "WRONG SIZE", RED
    if r["sim"] >= 0.95: return "VERY CLOSE", GREEN
    if r["sim"] >= 0.88: return "CLOSE", AMBER
    return "DIFFERENT", RED
prs = Presentation(); prs.slide_width, prs.slide_height = SW, SH
blank = prs.slide_layouts[6]
s = prs.slides.add_slide(blank); bg(s, INK)
tb(s, Inches(0.8), Inches(2.8), Inches(11.5), Inches(1.2), TITLE.upper(), 54, True, WHITE, "Cambria")
tb(s, Inches(0.8), Inches(4.1), Inches(11.5), Inches(0.8), "Designer vs Automated", 28, False, RGBColor(0xF4, 0xB4, 0xB6))
totals = {}
if len(sys.argv) > 5:
    for sl in json.load(open(sys.argv[5], encoding="utf-8")).get("slides", []):
        s = prs.slides.add_slide(blank); title(s, sl["title"], sl.get("subtitle", ""))
        secs = sl["sections"]; colw = 12.1 / len(secs)
        for i, sec in enumerate(secs):
            tb(s, Inches(0.6) + i * Inches(colw + 0.1), Inches(1.7), Inches(colw), Inches(5.4),
               [f"## {sec['heading']}"] + [f"- {t}" for t in sec["items"]], 12)


def add_brand(brand, pairs):
    for r in pairs: r["v"], r["col"] = verdict(r)
    cnt = lambda k: sum(1 for r in pairs if r["v"] == k)
    name = (brand or TITLE).title()
    s = prs.slides.add_slide(blank); title(s, f"{name} - at a glance", f"{len(pairs)} boards: {NOTE}")
    stats = [(cnt("VERY CLOSE"), "very close", "pixel similarity 95 % or more", GREEN), (cnt("CLOSE"), "close", "88-95 % similar", AMBER),
             (cnt("DIFFERENT"), "different", "under 88 % similar - layout or text differs", RED), (cnt("WRONG SIZE"), "wrong proportions", "aspect ratio differs from the designer's by more than 5 %", RED)]
    for i, (n, lab, sub, col) in enumerate(stats):
        x = Inches(0.6) + i * Inches(3.1); card(s, x, Inches(1.8), Inches(2.9), Inches(2.4))
        tb(s, x + Inches(0.2), Inches(1.95), Inches(2.5), Inches(0.9), f"{n}/{len(pairs)}", 44, True, col, "Cambria")
        tb(s, x + Inches(0.2), Inches(2.95), Inches(2.5), Inches(0.5), lab.capitalize(), 15, True, INK)
        tb(s, x + Inches(0.2), Inches(3.4), Inches(2.5), Inches(0.8), sub, 11, False, GREY)
    sims = [r["sim"] for r in pairs if r.get("sim")]
    avg = f" (average {sum(sims)/len(sims):.2f})" if sims else ""
    tb(s, Inches(0.6), Inches(4.6), Inches(12.1), Inches(2.4), ["## How it is measured",
       f"Each automated picture is paired with the designer's JPEG of the same S.No, size and shop. Similarity = 1 - mean grey-level difference of both pictures at the same reduced size{avg}. It rates the whole picture, so a different shop name or a small shift barely moves it - look at the slides.",
       "Proportions are compared from the pixel size of each picture."], 13)
    checked = [r for r in pairs if r.get("name_issues") is not None]
    if checked:
        bad = sum(1 for r in checked if r["name_issues"])
        tb(s, Inches(0.6), Inches(6.55), Inches(12.1), Inches(0.5), f"Name check (a name off the page, on top of another name or under a picture): {len(checked) - bad} of {len(checked)} boards clean" + (f", {bad} flagged" if bad else ""), 13, True, RED if bad else GREEN)
    totals[name] = {k: cnt(k) for k in ("VERY CLOSE", "CLOSE", "DIFFERENT", "WRONG SIZE")}
    for r in pairs:
        s = prs.slides.add_slide(blank); bg(s)
        m = re.match(r"(\d+) - (.*?) - (.*?) - (.*)", r["name"])  # NOSONAR - bounded, human-entered strings (file names / emails); no ReDoS exposure, rewrite would risk parsing changes
        sno, size, typ, shop = m.groups() if m else ("", "", "", r["name"])
        pill(s, Inches(0.5), Inches(0.35), Inches(1.15), Inches(0.5), f"S.No {int(sno)}" if sno else name, RED, 15)
        tb(s, Inches(1.85), Inches(0.3), Inches(7.6), Inches(0.6), shop, 24 if len(shop) <= 30 else 16, True, INK, "Cambria", anchor=1)
        tb(s, Inches(9.6), Inches(0.35), Inches(3.2), Inches(0.5), f"{size} - {typ}", 15, False, GREY, align=3, anchor=1)
        pill(s, Inches(0.5), Inches(0.98), Inches(2.6), Inches(0.36), r["v"] + (f" ({r['sim']*100:.0f} %)" if r.get("sim") and r["d"] else ""), r["col"], 11)
        issues = r.get("name_issues")
        if issues:
            pill(s, Inches(3.25), Inches(0.98), Inches(9.55), Inches(0.36), "Name check: " + "; ".join(issues)[:110], RED, 10)
        elif issues == []:
            pill(s, Inches(3.25), Inches(0.98), Inches(2.6), Inches(0.36), "Name check: OK", GREEN, 11)
        top, bottom = Inches(1.55), Inches(7.2)
        if not r["a"] or not r.get("d"):
            tb(s, Inches(0.5), top, Inches(12), Inches(0.5), "No automated picture was produced for this board." if not r["a"] else "No designer JPEG was found for this board.", 14, False, GREY)
            if r["a"]: place(s, r["a"], Inches(0.5), top + Inches(0.6), Inches(12.3), bottom - top - Inches(0.7))
            continue
        if ratio(r["dsz"]) >= 1.25:
            h2 = (bottom - top - Inches(0.7)) // 2
            tb(s, Inches(0.5), top, Inches(3), Inches(0.3), "DESIGNER", 11, True, GREY); place(s, r["d"], Inches(0.5), top + Inches(0.3), Inches(12.3), h2 - Inches(0.3))
            tb(s, Inches(0.5), top + h2 + Inches(0.15), Inches(3), Inches(0.3), "AUTOMATED", 11, True, RED); place(s, r["a"], Inches(0.5), top + h2 + Inches(0.45), Inches(12.3), h2 - Inches(0.3))
        else:
            tb(s, Inches(0.5), top, Inches(3), Inches(0.3), "DESIGNER", 11, True, GREY); place(s, r["d"], Inches(0.5), top + Inches(0.3), Inches(6.0), bottom - top - Inches(0.4))
            tb(s, Inches(6.9), top, Inches(3), Inches(0.3), "AUTOMATED", 11, True, RED); place(s, r["a"], Inches(6.9), top + Inches(0.3), Inches(6.0), bottom - top - Inches(0.4))


for brand, pairs in groups.items():
    add_brand(brand, pairs)
prs.save(sys.argv[2]); print("saved", len(prs.slides), totals)
