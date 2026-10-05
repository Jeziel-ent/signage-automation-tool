"""Build the Agarpathi 'Designer vs Automated' comparison deck from the evaluation runs.
    python tools/agarpathi_build_deck.py <out.pptx>
Uses eval_known/results.json (production behaviour) and eval_unseen/results.json (size never made by the designer)."""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from app import agarpathi_template as T  # noqa: E402
from tools.agarpathi_eval import CLOSE, EXACT, verdict  # noqa: E402

RUN = ROOT / "dataset_analysis" / "agarpathi_fidelity"
RED, INK, GREY, GREEN, AMBER, WHITE = (RGBColor(0xE3, 0x1E, 0x24), RGBColor(0x1F, 0x1F, 0x24), RGBColor(0x6B, 0x6B, 0x76),
                                       RGBColor(0x1E, 0x8E, 0x4E), RGBColor(0xE0, 0x8A, 0x00), RGBColor(0xFF, 0xFF, 0xFF))
SW, SH = Inches(13.333), Inches(7.5)


def load(mode):
    p = RUN / f"eval_{mode}" / "results.json"
    return {r["file"]: r for r in json.loads(p.read_text(encoding="utf-8"))} if p.exists() else {}


def classify(row, irreproducible: bool) -> str:
    """EXACT / CLOSE / OFF / ART (designer used pictures the master does not have) / FAILED / n/a"""
    if row is None:
        return "n/a"
    if row.get("status") != "done" or not row.get("geometry"):
        return "FAILED"
    ok, _ = verdict(row["geometry"], EXACT)
    if ok:
        return "EXACT"
    ok, _ = verdict(row["geometry"], CLOSE)
    if ok:
        return "CLOSE"
    return "ART" if irreproducible else "OFF"


def tb(slide, x, y, w, h, text, size=14, bold=False, color=INK, font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
        r.font.color.rgb = color
    return box


def pill(slide, x, y, w, h, text, fill, size=13):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.name = Pt(size), True, "Calibri"
    r.font.color.rgb = WHITE
    return s


def jpeg(path, max_side=1500):
    im = Image.open(path).convert("RGB")
    im.thumbnail((max_side, max_side))
    b = io.BytesIO()
    im.save(b, "JPEG", quality=82)
    b.seek(0)
    return b, im.size


def place(slide, path, x, y, w, h):
    """Picture scaled to fit (x, y, w, h) without cropping, centred in it."""
    b, (iw, ih) = jpeg(path)
    k = min(w / iw, h / ih)
    pw, ph = int(iw * k), int(ih * k)
    slide.shapes.add_picture(b, x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)


def bg(slide, color=WHITE):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


COL = {"EXACT": GREEN, "CLOSE": AMBER, "OFF": RED, "ART": GREY, "FAILED": RED, "n/a": GREY}
LABEL = {"EXACT": "EXACT MATCH", "CLOSE": "CLOSE", "OFF": "OFF", "ART": "DESIGNER ART N/A", "FAILED": "FAILED", "n/a": "-"}


def main(out):
    known, unseen = load("known"), load("unseen")
    tm = {t["file"]: t for t in T.templates()}
    # which designer boards use pictures the master does not contain
    import tools.agarpathi_loo as L
    md = {"L": json.load(open(next(L.CACHE.glob("76_-_125*")), encoding="utf-8")),
          "P": json.load(open(next(L.CACHE.glob("76_-_36*")), encoding="utf-8"))}
    cache = {f: L._master_objs(md[f]) for f in md}
    irrep = {}
    for r in known.values():
        t = tm.get(r["safe"] + ".json")
        objs, _top, W, H = cache[r["fam"]]
        ids = T.identify(objs, W, H)
        irrep[r["file"]] = bool(t) and not T.compatible(t, T.master_aspects(objs, ids, r["fam"]))

    rows = sorted(known.values(), key=lambda r: (r["sno"], r["w"]))
    kc = {r["file"]: classify(r, irrep[r["file"]]) for r in rows}
    uc = {r["file"]: classify(unseen.get(r["file"]), irrep[r["file"]]) for r in rows}
    n = len(rows)

    def count(d, k):
        return sum(1 for v in d.values() if v == k)

    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    blank = prs.slide_layouts[6]

    # --- 1 title
    s = prs.slides.add_slide(blank)
    bg(s, INK)
    tb(s, Inches(0.8), Inches(2.3), Inches(11.5), Inches(1), "AGARPATHI  -  SIGNAGE AUTOMATION", 40, True, WHITE, "Cambria")
    tb(s, Inches(0.8), Inches(3.4), Inches(11.5), Inches(0.7), "Designer JPEGs vs Automated Generated  -  template engine", 26, False, RGBColor(0xF4, 0xB4, 0xB6))
    tb(s, Inches(0.8), Inches(4.3), Inches(11.5), Inches(1.2),
       [f"{n} boards compared, matched by S.No (16 - 76), measured against the designer's own CDR files",
        "Automated: dual-master Darshan templates (125x48 in landscape, 36x48 in portrait), real CorelDRAW engine"], 16, False, RGBColor(0xC9, 0xC9, 0xD0))

    # --- 2 at a glance
    s = prs.slides.add_slide(blank)
    bg(s)
    tb(s, Inches(0.6), Inches(0.4), Inches(12), Inches(0.8), "At a glance", 34, True, INK, "Cambria")
    ok_k = count(kc, "EXACT")
    stats = [(f"{ok_k}/{n}", "boards EXACT at known sizes", "every picture, badge, box and shop-name within 2 % of the page of the designer's CDR", GREEN),
             (f"{count(kc, 'EXACT') + count(kc, 'CLOSE')}/{n}", "EXACT or CLOSE at known sizes", "CLOSE = within 5 % of the page", AMBER),
             (f"{count(uc, 'EXACT')}/{n}", "EXACT for a size never made before", "same test with every board of that size hidden from the engine", GREEN),
             (f"{count(kc, 'ART')}", "boards need artwork the master lacks", "designer swapped in other pictures (flat black box, two-line Sugandha badge)", GREY)]
    for i, (big, lab, sub, col) in enumerate(stats):
        x = Inches(0.6) + i * Inches(3.1)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), Inches(2.9), Inches(2.5))
        card.fill.solid(); card.fill.fore_color.rgb = RGBColor(0xF6, 0xF6, 0xF8); card.line.fill.background()
        tb(s, x + Inches(0.2), Inches(1.65), Inches(2.5), Inches(0.9), big, 44, True, col, "Cambria")
        tb(s, x + Inches(0.2), Inches(2.65), Inches(2.5), Inches(0.5), lab, 15, True, INK)
        tb(s, x + Inches(0.2), Inches(3.2), Inches(2.5), Inches(0.8), sub, 11, False, GREY)
    tb(s, Inches(0.6), Inches(4.4), Inches(12), Inches(2.6), [
        "How the score works",
        "Each automated board is opened in CorelDRAW and every element - Sugandha badge, BLACK STONE badge, Darshan roof logo, product box, table/box composite, shop-name lines - is compared with the same element in the designer's CDR.",
        "EXACT = every element within 2 % of the page (position) and 3 % (size). CLOSE = 5 % / 6 %. Tamil text is compared by right edge and height: the designer's Tamil spelling differs from the transliteration.",
        "Known sizes = production behaviour (the designer's board of that size is in the template library). Unseen sizes = the same boards re-run with every board of the same size removed."], 13, False, INK)
    tb(s, Inches(0.6), Inches(4.4), Inches(12), Inches(0.35), "", 1)

    # --- 3 breakdown table
    s = prs.slides.add_slide(blank)
    bg(s)
    tb(s, Inches(0.6), Inches(0.4), Inches(12), Inches(0.8), "Result breakdown", 34, True, INK, "Cambria")
    heads = ["", "EXACT", "CLOSE", "OFF", "Designer art n/a", "Failed"]
    tbl = s.shapes.add_table(4, 6, Inches(0.6), Inches(1.5), Inches(12.1), Inches(2.4)).table
    for c, h in enumerate(heads):
        tbl.cell(0, c).text = h
    def fill_row(r, label, d, fam=None):
        sel = {f: v for f, v in d.items() if fam is None or next(x for x in rows if x["file"] == f)["fam"] == fam}
        vals = [label, *[str(sum(1 for v in sel.values() if v == k)) for k in ("EXACT", "CLOSE", "OFF", "ART", "FAILED")]]
        for c, v in enumerate(vals):
            tbl.cell(r, c).text = v
    fill_row(1, "Known sizes - all boards", kc)
    fill_row(2, "Unseen sizes - all boards", uc)
    tbl.cell(3, 0).text = "Landscape / portrait split (known)"
    land = {f: v for f, v in kc.items() if next(x for x in rows if x["file"] == f)["fam"] == "L"}
    port = {f: v for f, v in kc.items() if next(x for x in rows if x["file"] == f)["fam"] == "P"}
    tbl.cell(3, 1).text = f"landscape {sum(v == 'EXACT' for v in land.values())}/{len(land)} exact"
    tbl.cell(3, 2).text = f"portrait/square {sum(v == 'EXACT' for v in port.values())}/{len(port)} exact"
    for r in range(4):
        for c in range(6):
            for p in tbl.cell(r, c).text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(14); run.font.name = "Calibri"; run.font.bold = (r == 0 or c == 0)

    prs.save(out)
    return rows, kc, uc, known, unseen, prs


if __name__ == "__main__":
    main(sys.argv[1])
